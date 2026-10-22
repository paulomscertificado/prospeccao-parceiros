"""Importação (CSV e Excel) e exportação (CSV, Excel e PDF)."""

import csv
import io
from datetime import date, datetime

from . import config, util

# Nome da coluna na planilha (sem acento, minúsculo) → campo do sistema.
SINONIMOS_LEAD = {
    "empresa": "empresa", "nome da empresa": "empresa", "razao social": "empresa", "escritorio": "empresa",
    "nome fantasia": "empresa",
    "nome": "responsavel", "responsavel": "responsavel", "contato": "responsavel", "nome do responsavel": "responsavel",
    "cargo": "cargo",
    "cnpj": "cnpj",
    "telefone": "telefone", "fone": "telefone", "tel": "telefone",
    "whatsapp": "whatsapp", "celular": "whatsapp", "zap": "whatsapp",
    "email": "email", "e-mail": "email",
    "site": "site", "website": "site",
    "instagram": "instagram", "linkedin": "linkedin",
    "cidade": "cidade", "municipio": "cidade",
    "estado": "estado", "uf": "estado",
    "endereco": "endereco",
    "especialidade": "especialidade", "observacoes": "observacoes",
    "quantidade de clientes": "qtd_clientes_estimada", "clientes": "qtd_clientes_estimada",
    "funcionarios": "funcionarios",
}

SINONIMOS_CERTIFICADO = {
    "cliente": "cliente", "nome": "cliente", "empresa": "cliente", "razao social": "cliente",
    "cnpj": "cnpj", "cpf/cnpj": "cnpj", "documento": "cnpj",
    "tipo": "tipo", "tipo de certificado": "tipo", "certificado": "tipo", "produto": "tipo",
    "emissao": "data_emissao", "data de emissao": "data_emissao",
    "vencimento": "data_vencimento", "data de vencimento": "data_vencimento", "validade": "data_vencimento",
    "parceiro": "parceiro", "parceiro responsavel": "parceiro", "contador": "parceiro",
    "telefone": "telefone", "email": "email", "e-mail": "email",
}


class ErroPlanilha(ValueError):
    pass


def ler_linhas(nome_arquivo: str, conteudo: bytes) -> list[dict]:
    """Lê CSV (separado por ; ou ,) ou Excel (.xlsx) e devolve uma lista de dicionários."""
    nome = (nome_arquivo or "").lower()
    if nome.endswith(".xlsx") or nome.endswith(".xlsm"):
        from openpyxl import load_workbook

        try:
            planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
        except Exception:
            raise ErroPlanilha("Não consegui abrir o arquivo Excel. Salve como .xlsx e tente de novo.") from None
        linhas = list(planilha.iter_rows(values_only=True))
        if not linhas:
            return []
        cabecalho = [str(c or "").strip() for c in linhas[0]]
        return [{cabecalho[i]: v for i, v in enumerate(l) if i < len(cabecalho)} for l in linhas[1:]
                if any(v not in (None, "") for v in l)]
    if nome.endswith(".xls"):
        raise ErroPlanilha("Arquivos .xls antigos não são aceitos. Abra no Excel e salve como .xlsx ou .csv.")
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErroPlanilha("Não consegui ler o arquivo. Salve como CSV UTF-8 e tente de novo.")
    amostra = texto[:2000]
    separador = ";" if amostra.count(";") >= amostra.count(",") else ","
    return [l for l in csv.DictReader(io.StringIO(texto), delimiter=separador)
            if any((v or "").strip() for v in l.values() if isinstance(v, str))]


def mapear(linha: dict, sinonimos: dict) -> dict:
    dados = {}
    for coluna, valor in linha.items():
        campo = sinonimos.get(util.chave_texto(str(coluna or "")))
        if campo and valor not in (None, ""):
            dados[campo] = valor if isinstance(valor, (date, datetime)) else str(valor).strip()
    return dados


def validar_lead(dados: dict) -> list[str]:
    problemas = []
    if not dados.get("empresa"):
        problemas.append("falta o nome da empresa")
    if dados.get("cnpj") and not util.cnpj_valido(dados["cnpj"]):
        problemas.append(f"CNPJ inválido ({dados['cnpj']})")
    if dados.get("email") and not util.email_valido(dados["email"]):
        problemas.append(f"e-mail inválido ({dados['email']})")
    for campo in ("telefone", "whatsapp"):
        if dados.get(campo) and not util.normalizar_telefone(dados[campo]):
            problemas.append(f"{campo} inválido ({dados[campo]})")
    if dados.get("estado") and dados["estado"].strip().upper() not in config.UFS:
        problemas.append(f"UF inválida ({dados['estado']})")
    return problemas


# ---------------------------------------------------------------- Exportação

def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, (date, datetime)):
        return util.data_br(valor)
    if isinstance(valor, float):
        return f"{valor:.2f}".replace(".", ",")
    return str(valor)


def gerar_csv(colunas: list[str], linhas: list[list]) -> bytes:
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";")
    escritor.writerow(colunas)
    for linha in linhas:
        escritor.writerow([_texto(v) for v in linha])
    return ("﻿" + saida.getvalue()).encode("utf-8")  # BOM: o Excel abre com acentos certos


def gerar_xlsx(titulo: str, colunas: list[str], linhas: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    livro = Workbook()
    aba = livro.active
    aba.title = titulo[:31]
    aba.append(colunas)
    for celula in aba[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F3A64")
    for linha in linhas:
        aba.append([v if isinstance(v, (int, float, date, datetime)) and not isinstance(v, bool) else _texto(v)
                    for v in linha])
    for i, coluna in enumerate(colunas, start=1):
        largura = max([len(coluna)] + [len(_texto(l[i - 1])) for l in linhas[:200]])
        aba.column_dimensions[aba.cell(row=1, column=i).column_letter].width = min(max(10, largura + 2), 50)
    aba.freeze_panes = "A2"
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _latin(texto: str) -> str:
    trocas = {"→": "->", "–": "-", "—": "-", "“": '"', "”": '"', "‘": "'", "’": "'", "…": "...", "•": "-"}
    for a, b in trocas.items():
        texto = texto.replace(a, b)
    return texto.encode("latin-1", "replace").decode("latin-1")


def gerar_pdf(titulo: str, secoes: list[tuple[str, list[str], list[list]]], subtitulo: str = "") -> bytes:
    """secoes: [(título da seção, colunas, linhas)]. Página deitada (paisagem)."""
    from fpdf import FPDF, FontFace

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 15)
    pdf.set_text_color(15, 58, 100)
    pdf.cell(0, 9, _latin(titulo), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(90, 100, 110)
    pdf.cell(0, 5, _latin(subtitulo or f"MS Certificação Digital · gerado em {datetime.now():%d/%m/%Y %H:%M}"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    largura_util = pdf.w - pdf.l_margin - pdf.r_margin
    for nome_secao, colunas, linhas in secoes:
        if nome_secao:
            pdf.set_font("Helvetica", "B", 11)
            pdf.set_text_color(15, 58, 100)
            pdf.cell(0, 7, _latin(nome_secao), new_x="LMARGIN", new_y="NEXT")
        if not colunas:
            continue
        pdf.set_text_color(30, 30, 30)
        pdf.set_font("Helvetica", "", 8)
        with pdf.table(width=largura_util, line_height=5, text_align="LEFT",
                       headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255),
                                                                  fill_color=(15, 58, 100))) as tabela:
            cab = tabela.row()
            for c in colunas:
                cab.cell(_latin(c))
            for linha in linhas:
                r = tabela.row()
                for v in linha:
                    r.cell(_latin(_texto(v))[:120])
        pdf.ln(4)
    return bytes(pdf.output())
